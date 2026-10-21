#!/usr/bin/env python3
"""
scripts/generate_template.py
Generates an executive-ready, beautifully styled Excel test case template for the QA Agent System.
Uses Segoe UI, modern corporate navy palettes, pastel pill badge styling for severity/readiness,
zebra striping, custom row heights, and auto-filters.
"""

import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# Typography
FONT_NAME = "Segoe UI"
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
DATA_FONT = Font(name=FONT_NAME, size=10, color="1E293B")
BOLD_DATA_FONT = Font(name=FONT_NAME, size=10, bold=True, color="1E293B")
NOTE_FONT = Font(name=FONT_NAME, size=9, italic=True, color="64748B")

# Color Palette (Corporate Modern Navy & Slate)
PRIMARY_NAVY = "0F294A"    # Deep corporate navy for Sheet 1
SECONDARY_NAVY = "1E3A8A"  # Indigo navy for Sheet 2
ACCENT_SLATE = "334155"    # Dark slate

FILL_HEADER_1 = PatternFill(start_color=PRIMARY_NAVY, end_color=PRIMARY_NAVY, fill_type="solid")
FILL_HEADER_2 = PatternFill(start_color=SECONDARY_NAVY, end_color=SECONDARY_NAVY, fill_type="solid")
FILL_ZEBRA_EVEN = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
FILL_ZEBRA_ODD = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

# Pastel Badges for Severity
FILL_CHAN = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")  # Soft Rose Red
FONT_CHAN = Font(name=FONT_NAME, size=10, bold=True, color="991B1B")

FILL_SAI = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")   # Soft Amber Yellow
FONT_SAI = Font(name=FONT_NAME, size=10, bold=True, color="92400E")

FILL_NHO = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")   # Soft Sky Blue
FONT_NHO = Font(name=FONT_NAME, size=10, bold=True, color="075985")

# Pastel Badges for Automation Ready
FILL_READY_YES = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # Soft Emerald
FONT_READY_YES = Font(name=FONT_NAME, size=10, bold=True, color="166534")

FILL_READY_NO = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")  # Neutral Slate
FONT_READY_NO = Font(name=FONT_NAME, size=10, color="475569")

# Borders
BORDER_SUBTLE = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0')
)
BORDER_HEADER_BOTTOM = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='medium', color='0F172A')
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_severity_cell(cell):
    val = str(cell.value or "").strip().lower()
    if "chặn" in val:
        cell.fill = FILL_CHAN
        cell.font = FONT_CHAN
    elif "sai" in val:
        cell.fill = FILL_SAI
        cell.font = FONT_SAI
    elif "nhỏ" in val:
        cell.fill = FILL_NHO
        cell.font = FONT_NHO
    cell.alignment = ALIGN_CENTER

def style_ready_cell(cell):
    val = str(cell.value or "").strip().lower()
    if val == "yes":
        cell.fill = FILL_READY_YES
        cell.font = FONT_READY_YES
    else:
        cell.fill = FILL_READY_NO
        cell.font = FONT_READY_NO
    cell.alignment = ALIGN_CENTER

def style_status_cell(cell):
    val = str(cell.value or "").strip().lower()
    if val == "pass":
        cell.fill = FILL_READY_YES
        cell.font = FONT_READY_YES
    elif val == "fail":
        cell.fill = FILL_CHAN
        cell.font = FONT_CHAN
    elif val == "skip":
        cell.fill = FILL_READY_NO
        cell.font = FONT_READY_NO
    else:  # Mặc định: Pending
        cell.fill = FILL_SAI
        cell.font = FONT_SAI
    cell.alignment = ALIGN_CENTER

def build_template(output_path: Path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    # ====================================================
    # Sheet 1: Test Cases
    # ====================================================
    ws_tc = wb.create_sheet(title="Test Cases")
    ws_tc.views.sheetView[0].showGridLines = True
    ws_tc.row_dimensions[1].height = 32

    tc_headers = [
        "Test Case ID",
        "Module",
        "Test Case Name",
        "Preconditions",
        "Test Data",
        "Test Steps",
        "Expected Result",
        "Severity",
        "Automation Ready",
        "Automation Blocker",
        "Requirement Reference",
        "Notes",
        "Status"
    ]
    ws_tc.append(tc_headers)

    # Dropdowns
    dv_severity = DataValidation(type="list", formula1='"Chặn,Sai,Nhỏ"', allow_blank=True)
    dv_severity.error = "Vui lòng chọn giá trị hợp lệ: Chặn, Sai, hoặc Nhỏ"
    dv_severity.errorTitle = "Giá trị không hợp lệ"
    ws_tc.add_data_validation(dv_severity)
    dv_severity.add("H2:H500")

    dv_auto = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    dv_auto.error = "Vui lòng chọn Yes hoặc No"
    dv_auto.errorTitle = "Giá trị không hợp lệ"
    ws_tc.add_data_validation(dv_auto)
    dv_auto.add("I2:I500")

    dv_status = DataValidation(type="list", formula1='"Pass,Fail,Pending,Skip"', allow_blank=True)
    dv_status.error = "Vui lòng chọn trạng thái: Pass, Fail, Pending, hoặc Skip"
    dv_status.errorTitle = "Trạng thái không hợp lệ"
    ws_tc.add_data_validation(dv_status)
    dv_status.add("M2:M500")

    # Sample rows showcasing Chặn and Sai with default Pending / Pass status
    sample_rows = [
        [
            "TC-AUTH-001",
            "Authentication",
            "Đăng nhập thành công với thông tin hợp lệ (Happy Path)",
            "Tài khoản đã kích hoạt, không có cảnh báo máy lạ",
            "Tên đăng nhập: user01, Mật khẩu: Pass@123",
            "1. Mở trang /dang-nhap\n2. Nhập username và password\n3. Bấm Đăng nhập",
            "Hệ thống xác thực thành công và điều hướng vào Dashboard",
            "Chặn",
            "Yes",
            "",
            "SPEC-AU-01 / BR§2.1",
            "Luồng đăng nhập cốt lõi sống còn",
            "Pass"
        ],
        [
            "TC-AUTH-002",
            "Authentication",
            "Chống spam click nút Đăng nhập khi mạng lag",
            "Nhập đúng thông tin tài khoản",
            "Tên đăng nhập: user01, Mật khẩu: Pass@123",
            "1. Nhập thông tin đăng nhập\n2. Bấm liên tiếp 3 lần nút Đăng nhập trong 500ms",
            "Nút chuyển trạng thái loading, vô hiệu hoá click trùng, chỉ gửi 1 request API duy nhất",
            "Chặn",
            "Yes",
            "",
            "SPEC-HT-26 (useSubmitOnce)",
            "Bảo vệ chống nhảy bộ đếm sai khoá tài khoản oan",
            "Pass"
        ],
        [
            "TC-AUTH-003",
            "Authentication",
            "Chuyển đổi Tab định danh (Mã NV -> SĐT) tự động xoá ô nhập",
            "Đang ở màn hình đăng nhập mặc định Tab Mã NV",
            "Chuỗi gõ dở dang: NV009988",
            "1. Gõ chuỗi 'NV009988' vào ô Mã NV\n2. Bấm đổi sang Tab 'Số điện thoại'",
            "Giao diện chuyển sang Tab SĐT, ô định danh bị xoá trắng chuỗi cũ, placeholder đổi sang dạng SĐT",
            "Sai",
            "Yes",
            "",
            "SPEC-AU-49 / VBO-DN-09",
            "Tránh người dùng nhầm lẫn mã NV dưới nhãn SĐT",
            "Pending"
        ]
    ]

    for r in sample_rows:
        ws_tc.append(r)

    # Formatting Sheet 1
    ws_tc.freeze_panes = "A2"
    ws_tc.auto_filter.ref = f"A1:{get_column_letter(len(tc_headers))}500"

    for cell in ws_tc[1]:
        cell.font = HEADER_FONT
        cell.fill = FILL_HEADER_1
        cell.alignment = ALIGN_HEADER
        cell.border = BORDER_HEADER_BOTTOM

    for row_idx, row in enumerate(ws_tc.iter_rows(min_row=2, max_row=len(sample_rows) + 1, min_col=1, max_col=len(tc_headers)), start=2):
        ws_tc.row_dimensions[row_idx].height = 28
        is_odd = (row_idx % 2 == 1)
        row_fill = FILL_ZEBRA_ODD if is_odd else FILL_ZEBRA_EVEN

        for cell in row:
            cell.font = DATA_FONT
            cell.border = BORDER_SUBTLE
            cell.fill = row_fill

            if cell.column == 1: # ID
                cell.font = BOLD_DATA_FONT
                cell.alignment = ALIGN_CENTER
            elif cell.column == 8: # Severity
                style_severity_cell(cell)
            elif cell.column == 9: # Automation Ready
                style_ready_cell(cell)
            elif cell.column == 12: # Notes
                cell.font = NOTE_FONT
                cell.alignment = ALIGN_LEFT
            elif cell.column == 13: # Status
                style_status_cell(cell)
            else:
                cell.alignment = ALIGN_LEFT

    col_widths_tc = {
        "A": 16, "B": 18, "C": 32, "D": 28, "E": 28, "F": 38,
        "G": 38, "H": 15, "I": 18, "J": 26, "K": 22, "L": 26, "M": 16
    }
    for col_letter, width in col_widths_tc.items():
        ws_tc.column_dimensions[col_letter].width = width

    # ====================================================
    # Sheet 2: Requirement Gaps
    # ====================================================
    ws_gaps = wb.create_sheet(title="Requirement Gaps")
    ws_gaps.views.sheetView[0].showGridLines = True
    ws_gaps.row_dimensions[1].height = 32

    gap_headers = [
        "Issue ID", "Related Flow", "Current Evidence", "Uncertainty",
        "Business Impact", "Question for BA/PO", "Suggested Options", "Status"
    ]
    ws_gaps.append(gap_headers)

    dv_gap_status = DataValidation(type="list", formula1='"Pending,Answered,Skipped"', allow_blank=True)
    ws_gaps.add_data_validation(dv_gap_status)
    dv_gap_status.add("H2:H200")

    sample_gaps = [
        [
            "GAP-001",
            "Khóa tài khoản khi nhập sai",
            "Code có LOGIN_THROTTLE.MAX_ATTEMPTS = 5",
            "Chưa rõ thời gian khoá bao lâu (15 phút hay vĩnh viễn cần admin mở)",
            "Ảnh hưởng đến thiết kế test timeout và assertion thông báo",
            "Tài khoản bị khoá sau 5 lần nhập sai sẽ tự mở sau bao lâu?",
            "A: Khóa 15 phút tự mở; B: Cần Admin reset thủ công",
            "Pending"
        ],
        [
            "GAP-002",
            "Cảnh báo đăng nhập máy lạ",
            "MayLaNotice xuất hiện khi mayLa === true",
            "Có cần gửi email/SMS thông báo khi phát hiện máy lạ hay không?",
            "Ảnh hưởng đến kịch bản kiểm tra tích hợp gửi thông báo",
            "Hệ thống chỉ hiện popup cảnh báo trên UI hay có bắn email cảnh báo?",
            "A: Chỉ hiện popup cảnh báo, không gửi mail; B: Bắn email thông báo",
            "Answered"
        ]
    ]

    for g in sample_gaps:
        ws_gaps.append(g)

    ws_gaps.freeze_panes = "A2"
    ws_gaps.auto_filter.ref = f"A1:{get_column_letter(len(gap_headers))}200"

    for cell in ws_gaps[1]:
        cell.font = HEADER_FONT
        cell.fill = FILL_HEADER_2
        cell.alignment = ALIGN_HEADER
        cell.border = BORDER_HEADER_BOTTOM

    for row_idx, row in enumerate(ws_gaps.iter_rows(min_row=2, max_row=len(sample_gaps) + 1, min_col=1, max_col=len(gap_headers)), start=2):
        ws_gaps.row_dimensions[row_idx].height = 28
        row_fill = FILL_ZEBRA_ODD if (row_idx % 2 == 1) else FILL_ZEBRA_EVEN

        for cell in row:
            cell.font = DATA_FONT
            cell.border = BORDER_SUBTLE
            cell.fill = row_fill

            if cell.column == 1:
                cell.font = BOLD_DATA_FONT
                cell.alignment = ALIGN_CENTER
            elif cell.column == 8: # Status
                st = str(cell.value or "").strip().lower()
                if st == "answered":
                    cell.fill = FILL_READY_YES
                    cell.font = FONT_READY_YES
                elif st == "pending":
                    cell.fill = FILL_SAI
                    cell.font = FONT_SAI
                else:
                    cell.fill = FILL_READY_NO
                    cell.font = FONT_READY_NO
                cell.alignment = ALIGN_CENTER
            else:
                cell.alignment = ALIGN_LEFT

    col_widths_gap = {
        "A": 14, "B": 24, "C": 28, "D": 32, "E": 28, "F": 36, "G": 32, "H": 16
    }
    for col_letter, width in col_widths_gap.items():
        ws_gaps.column_dimensions[col_letter].width = width

    # ====================================================
    # Sheet 3: Coverage Summary (Executive KPI Card View)
    # ====================================================
    ws_sum = wb.create_sheet(title="Coverage Summary")
    ws_sum.views.sheetView[0].showGridLines = True
    ws_sum.column_dimensions["A"].width = 28
    ws_sum.column_dimensions["B"].width = 65

    # Title Banner
    ws_sum.row_dimensions[1].height = 40
    ws_sum.append(["BÁO CÁO TỔNG HỢP ĐỘ BAO PHỦ VÀ RỦI RO KIỂM THỬ", ""])
    ws_sum.merge_cells("A1:B1")
    ws_sum["A1"].font = TITLE_FONT
    ws_sum["A1"].fill = FILL_HEADER_1
    ws_sum["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    summary_items = [
        ("Tên Phân hệ / Module:", "Authentication / Đăng nhập"),
        ("Tổng số Test Cases Thiết kế:", "3 (Mẫu - Tiêu chuẩn thực tế: 7–10 Chặn, 5–7 Sai)"),
        ("Số ca Sẵn sàng Tự động hoá (Yes):", "3 (100%)"),
        ("Số ca Yêu cầu Kiểm thử Thủ công (No):", "0"),
        ("Phân bổ Mức độ Nghiêm trọng:", "Chặn (Nặng): 2 | Sai: 1 | Nhỏ: 0"),
        ("Tỷ lệ Truy vết Yêu cầu (Traceability Rate):", "100% (3/3 yêu cầu trọng yếu đã có test case bao phủ)"),
        ("Yêu cầu Trọng yếu Chưa có Test Case (Orphan):", "Không có (0) - Đạt chuẩn 100% bao phủ nghiệp vụ"),
        ("Test Case Ảo Không có Căn cứ Code/Spec:", "Không có (0) - 100% test case có mã đối soát thực tế"),
        ("Test Case Trùng lặp Mục tiêu (Duplicates):", "Không có (0) - Bộ test case độc lập và tối giản"),
        ("Các Luồng Nghiệp vụ Đã Bao Phủ:", "• Luồng Đăng nhập hợp lệ chính thống (Happy Path)\n• Cơ chế phòng thủ chống spam click submit liên tục (useSubmitOnce)\n• Chuyển đổi qua lại giữa Tab Mã NV và Số điện thoại"),
        ("Vùng Rủi ro Chưa Bao Phủ / Tồn đọng:", "• Luồng xác thực đăng nhập một lần (SSO OAuth2 qua Google/Microsoft)\n• Luồng mở khoá bằng mã PIN vật lý bảo mật phần cứng"),
        ("Giả định Môi trường & Dữ liệu:", "• Môi trường Staging không bật Google reCAPTCHA v3\n• Database có sẵn tài khoản mẫu đã kích hoạt và không bị khoá")
    ]

    for label, val in summary_items:
        ws_sum.append([label, val])

    for row_idx in range(2, len(summary_items) + 2):
        ws_sum.row_dimensions[row_idx].height = 30
        cell_lbl = ws_sum.cell(row=row_idx, column=1)
        cell_val = ws_sum.cell(row=row_idx, column=2)

        cell_lbl.font = BOLD_DATA_FONT
        cell_lbl.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        cell_lbl.border = BORDER_SUBTLE
        cell_lbl.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        cell_val.font = DATA_FONT
        cell_val.fill = FILL_ZEBRA_EVEN
        cell_val.border = BORDER_SUBTLE
        cell_val.alignment = ALIGN_LEFT

    # Save
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"[OK] Generated high-end polished Excel template at: {output_path}")

if __name__ == "__main__":
    out = Path(__file__).resolve().parent.parent / "templates" / "testcase-template.xlsx"
    if len(sys.argv) > 1:
        out = Path(sys.argv[1])
    build_template(out)
