#!/usr/bin/env python3
"""
scripts/read_approved_testcases.py
Reads an approved Excel test case file, validates its schema,
and selects strictly rows where `Automation Ready = Yes`.
"""

import sys
import json
from pathlib import Path
from typing import Dict, Any, List

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl

REQUIRED_COLUMNS = [
    "Test Case ID", "Module", "Test Case Name", "Preconditions",
    "Test Data", "Test Steps", "Expected Result", "Severity", "Automation Ready"
]

def parse_approved_workbook(excel_path: Path) -> Dict[str, Any]:
    if not excel_path.exists():
        return {"success": False, "error": f"File not found: {excel_path}"}

    try:
        wb = openpyxl.load_workbook(excel_path, data_only=True)
    except Exception as e:
        return {"success": False, "error": f"Failed to open workbook: {str(e)}"}

    if "Test Cases" not in wb.sheetnames:
        return {
            "success": False,
            "error": f"Missing required sheet 'Test Cases'. Available sheets: {wb.sheetnames}"
        }

    ws = wb["Test Cases"]
    header_row = [cell.value for cell in ws[1]]
    col_map = {}
    for idx, name in enumerate(header_row):
        if name:
            col_map[str(name).strip()] = idx + 1

    missing = [c for c in REQUIRED_COLUMNS if c not in col_map]
    if missing:
        return {
            "success": False,
            "error": f"Missing required columns in 'Test Cases' sheet: {missing}"
        }

    all_cases = []
    automation_ready_cases = []
    skipped_cases = []

    for row_idx in range(2, ws.max_row + 1):
        tc_id = ws.cell(row=row_idx, column=col_map["Test Case ID"]).value
        if not tc_id:
            continue

        tc_id_str = str(tc_id).strip()
        auto_ready_val = ws.cell(row=row_idx, column=col_map["Automation Ready"]).value
        auto_ready_str = str(auto_ready_val).strip() if auto_ready_val else "No"
        is_ready = auto_ready_str.lower() == "yes"

        tc_item = {
            "row_number": row_idx,
            "id": tc_id_str,
            "module": str(ws.cell(row=row_idx, column=col_map["Module"]).value or "").strip(),
            "name": str(ws.cell(row=row_idx, column=col_map["Test Case Name"]).value or "").strip(),
            "preconditions": str(ws.cell(row=row_idx, column=col_map["Preconditions"]).value or "").strip(),
            "test_data": str(ws.cell(row=row_idx, column=col_map["Test Data"]).value or "").strip(),
            "steps": str(ws.cell(row=row_idx, column=col_map["Test Steps"]).value or "").strip(),
            "expected_result": str(ws.cell(row=row_idx, column=col_map["Expected Result"]).value or "").strip(),
            "severity": str(ws.cell(row=row_idx, column=col_map["Severity"]).value or "").strip(),
            "automation_ready": "Yes" if is_ready else "No",
            "automation_blocker": str(ws.cell(row=row_idx, column=col_map.get("Automation Blocker", 10)).value or "").strip() if "Automation Blocker" in col_map else "",
            "requirement_ref": str(ws.cell(row=row_idx, column=col_map.get("Requirement Reference", 11)).value or "").strip() if "Requirement Reference" in col_map else "",
            "notes": str(ws.cell(row=row_idx, column=col_map.get("Notes", 12)).value or "").strip() if "Notes" in col_map else "",
            "status": str(ws.cell(row=row_idx, column=col_map.get("Status", 13)).value or "Pending").strip() if "Status" in col_map else "Pending"
        }

        all_cases.append(tc_item)
        if is_ready:
            automation_ready_cases.append(tc_item)
        else:
            skipped_cases.append(tc_item)

    return {
        "success": True,
        "excel_path": str(excel_path.resolve()),
        "total_cases": len(all_cases),
        "automation_ready_count": len(automation_ready_cases),
        "skipped_count": len(skipped_cases),
        "ready_testcases": automation_ready_cases,
        "skipped_testcases": skipped_cases
    }

def print_summary(res: Dict[str, Any]):
    print(f"[OK] Ingested: {Path(res['excel_path']).name}")
    print(f"Total: {res['total_cases']} | Ready: {res['automation_ready_count']} | Skipped: {res['skipped_count']}")
    if res["ready_testcases"]:
        print("\nAutomation Ready Test Cases:")
        for tc in res["ready_testcases"]:
            print(f"  - [{tc['id']}] {tc['name']} ({tc['severity']})")
    if res["skipped_testcases"]:
        print("\nSkipped Test Cases (Manual/Blocked):")
        for tc in res["skipped_testcases"]:
            blocker = f" - Reason: {tc['automation_blocker']}" if tc.get("automation_blocker") else ""
            print(f"  - [{tc['id']}] {tc['name']}{blocker}")

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="Read and validate approved Excel test case workbook.")
    parser.add_argument("workbook", help="Path to approved Excel workbook (.xlsx)")
    parser.add_argument("-o", "--output", help="Write full JSON to file instead of stdout (saves LLM tokens)")
    parser.add_argument("--summary", action="store_true", help="Print only executive summary table")
    parser.add_argument("--ready-only", action="store_true", help="Output only automation-ready test cases")
    parser.add_argument("--id", dest="tc_id", help="Extract a single test case by ID")
    parser.add_argument("--json", action="store_true", help="Output raw JSON to stdout")

    args = parser.parse_args()

    file_p = Path(args.workbook)
    res = parse_approved_workbook(file_p)

    if not res.get("success"):
        print(f"[ERROR] {res.get('error')}", file=sys.stderr)
        sys.exit(1)

    # 1. If output file requested, write full JSON to disk
    if args.output:
        out_p = Path(args.output)
        out_p.parent.mkdir(parents=True, exist_ok=True)
        with open(out_p, "w", encoding="utf-8") as f:
            json.dump(res, f, indent=2, ensure_ascii=False)
        print(f"[SAVED] Full test cases saved to: {out_p}")
        print_summary(res)
        sys.exit(0)

    # 2. If single ID requested
    if args.tc_id:
        target = next((tc for tc in res["ready_testcases"] + res["skipped_testcases"] if tc["id"].lower() == args.tc_id.lower()), None)
        if target:
            print(json.dumps(target, indent=2, ensure_ascii=False))
        else:
            print(f"[ERROR] Test case '{args.tc_id}' not found.", file=sys.stderr)
            sys.exit(1)
        sys.exit(0)

    # 3. If ready-only requested
    if args.ready_only:
        payload = {
            "success": True,
            "total_ready": len(res["ready_testcases"]),
            "ready_testcases": res["ready_testcases"]
        }
        print(json.dumps(payload, indent=2 if args.json else None, ensure_ascii=False))
        sys.exit(0)

    # 4. If raw JSON requested
    if args.json:
        print(json.dumps(res, indent=2, ensure_ascii=False))
        sys.exit(0)

    # 5. Default / --summary: Print high-signal token-optimized summary
    print_summary(res)
