"""Golden-project integration test using real filesystem, discovery, Excel, and RTM engines."""

import sqlite3
import sys
from pathlib import Path

# Ensure project root is in sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from openpyxl import load_workbook

from scripts.codegraph_discovery import (
    MethodDiscoveryEngine,
    JavaScriptTypeScriptASTAdapter,
)
from scripts.export_testcases import export_to_excel
from scripts.read_approved_testcases import parse_approved_workbook
from scripts.state_manager import WorkflowStateManager
from scripts.traceability_validator import validate_traceability


FIXTURE = Path(__file__).parent / "fixtures" / "golden-sample-project"


def test_golden_pipeline_without_mocks(tmp_path):
    spec = FIXTURE / "spec.md"
    assert "SPEC-AUTH-001" in spec.read_text(encoding="utf-8")

    state = WorkflowStateManager(tmp_path, module_name="golden-login")
    initialized = state.init_state(
        "golden-login",
        {"root": str(FIXTURE), "source": "golden-project"},
        [str(spec)],
    )
    assert initialized["current_state"] == "INITIALIZED"

    # Step 1: Context discovery
    engine = MethodDiscoveryEngine(FIXTURE)
    discovery = engine.discover_and_evaluate("login")
    engine.close()
    assert discovery["status"] == "FOUND"
    assert discovery["candidate"]["name"] == "login"
    assert discovery["discovery_source"] == "fallback_ast_scan"
    state.transition("CONTEXT_DISCOVERY", note="Completed AST context discovery")

    # Step 2: Gap analysis
    state.transition("GAP_ANALYSIS", note="Evaluated spec vs codebase, zero blocking gaps")

    # Step 3: Testcase generation
    payload = {
        "module": "Golden Login",
        "requirements": ["SPEC-AUTH-001"],
        "testcases": [{
            "id": "TC-AUTH-001",
            "name": "Đăng nhập bằng thông tin hợp lệ",
            "severity": "Chặn",
            "priority": "P0",
            "preconditions": "Tài khoản đang hoạt động",
            "test_data": "Tài khoản kiểm thử hợp lệ",
            "steps": "1. Mở màn hình đăng nhập\n2. Nhập thông tin hợp lệ\n3. Bấm Đăng nhập",
            "expected_result": "Màn hình bảng điều khiển được hiển thị",
            "requirement_ref": "SPEC-AUTH-001",
            "automation_ready": "Yes",
            "code_reference": f"{discovery['candidate']['file_path']}#{discovery['candidate']['name']}",
        }],
    }
    state.transition("TESTCASE_GENERATION", note="Designed risk-based test suite")

    # Step 4: Quality review
    audit = validate_traceability(payload)
    assert audit["valid"] is True
    assert audit["metrics"]["coverage_percentage"] == 100.0
    state.transition("TESTCASE_REVIEW", note="Internal quality review and RTM audit passed")

    # Step 5: Export to Excel and wait for approval
    workbook = tmp_path / "golden-login-testcases.xlsx"
    export_to_excel(payload, workbook)
    assert workbook.exists() and workbook.stat().st_size > 0
    loaded = load_workbook(workbook, read_only=True)
    assert "Test Cases" in loaded.sheetnames
    loaded.close()
    state.transition(
        "WAITING_FOR_APPROVED_TESTCASES",
        note="Exported workbook awaiting QA Lead signoff",
        metadata={"artifacts": {"approved_excel": str(workbook)}},
    )

    # Step 6: Round-Trip Ingestion verification (reading back approved test cases)
    state.transition(
        "AUTOMATION_IMPLEMENTATION",
        note="Approved by QA Lead, verified round-trip",
        approval_token="qa-lead-approved",
    )
    approved_result = parse_approved_workbook(workbook)
    assert approved_result["success"] is True
    assert approved_result["total_cases"] == 1
    assert approved_result["automation_ready_count"] == 1
    assert approved_result["ready_testcases"][0]["id"] == "TC-AUTH-001"
    assert approved_result["ready_testcases"][0]["name"] == "Đăng nhập bằng thông tin hợp lệ"

    state.transition("COMPLETED", note="Golden testcase roundtrip pipeline verified successfully")

    final_state = state.load()
    assert final_state["current_state"] == "COMPLETED"
    assert "WAITING_FOR_APPROVED_TESTCASES" in final_state["completed_steps"]
    assert len(final_state["history"]) >= 6


def test_polyglot_adapter_and_truncation_observability(tmp_path):
    # 1. Verify JavaScriptTypeScriptASTAdapter directly executes without NotImplementedError
    adapter = JavaScriptTypeScriptASTAdapter(synonyms={"login": ["login", "signin"]})
    candidates = adapter.discover(FIXTURE, "login")
    assert len(candidates) > 0
    assert any(c["name"] == "login" for c in candidates)

    # 2. Verify MethodDiscoveryEngine wires adapter seamlessly
    engine = MethodDiscoveryEngine(FIXTURE)
    assert len(engine.adapters) >= 1
    assert isinstance(engine.adapters[0], JavaScriptTypeScriptASTAdapter)

    # 3. Verify SQLite candidate truncation transparency
    db_dir = tmp_path / ".codegraph"
    db_dir.mkdir(parents=True, exist_ok=True)
    db_file = db_dir / "codegraph.db"
    conn = sqlite3.connect(db_file)
    conn.execute("""
        CREATE TABLE nodes (
            id TEXT PRIMARY KEY,
            kind TEXT,
            name TEXT,
            qualified_name TEXT,
            file_path TEXT,
            start_line INTEGER,
            end_line INTEGER,
            signature TEXT,
            docstring TEXT,
            is_exported INTEGER
        )
    """)
    # Insert 3 matching methods
    for i in range(1, 4):
        conn.execute(
            "INSERT INTO nodes VALUES (?, 'method', 'handlePayment', 'Order::handlePayment', 'order.ts', ?, ?, '() => void', '', 1)",
            (f"node_{i}", i * 10, i * 10 + 5)
        )
    conn.commit()
    conn.close()

    mock_engine = MethodDiscoveryEngine(tmp_path)
    # Request max_candidates=2 when 3 exist
    results = mock_engine.discover_candidates("handlePayment", max_candidates=2)
    assert len(results) == 2
    # Verify transparent truncation warning was recorded
    warnings = mock_engine.last_traversal_report.get("warnings", [])
    assert any("truncated to top 2" in w for w in warnings)
    mock_engine.close()
    engine.close()
